"""Structural tests for the consolidated run_all workbook generator.

Builds a tiny synthetic dataset (two reference runs that share the prefix "ab"
then split on an ambiguous [c|d] event) and checks that run_all.build writes
exactly three workbooks per dataset -- one per algorithm -- each carrying the
Summary sheet plus the per-trace detail sheets, and that no CSV is produced.
"""

from __future__ import annotations

import json
from pathlib import Path

import openpyxl
import pytest

from evaluation import run_all

# A dataset name that exists in common.refs.CASE_KEY, so no registry change is
# needed; its case key is "routine_id" and its short name is "SH".
DATASET = "smart_home"
CASE_KEY = "routine_id"

LABELS = {"a": 0, "b": 1, "c": 2, "d": 3, "e": 4, "f": 5}
RUNS = ["ABCE", "ABDF"]  # chr(index + 'A'); shared prefix AB, split at index 2


def _event_single(name):
    return (f'\t\t<event>\n\t\t\t<string key="{CASE_KEY}" value="1" />\n'
            f'\t\t\t<string key="concept:name" value="{name}" />\n\t\t</event>')


def _event_ambiguous(names):
    lines = [f'\t\t<event>\n\t\t\t<string key="{CASE_KEY}" value="1" />',
             '\t\t\t<container key="uncertainty:discrete_weak">']
    for n in names:
        lines.append('\t\t\t\t<container key="uncertainty:entry">'
                     f'<string key="concept:name" value="{n}" /></container>')
    lines += ['\t\t\t</container>', '\t\t</event>']
    return "\n".join(lines)


def _log(events):
    return ('<?xml version="1.0" encoding="UTF-8" ?>\n'
            '<log xes.version="1849-2016" xmlns="http://www.xes-standard.org/">\n'
            '\t<trace>\n' + "\n".join(events) + "\n\t</trace>\n</log>\n")


@pytest.fixture
def data_dir(tmp_path):
    ddir = tmp_path / DATASET
    ddir.mkdir()
    (ddir / "labels.json").write_text(json.dumps(LABELS), encoding="utf-8")
    (ddir / "runs.txt").write_text("\n".join(RUNS) + "\n", encoding="utf-8")
    # true path is a b c e; the third event is observed ambiguously as [c|d]
    (ddir / "gold_standard.xes").write_text(
        _log([_event_single(x) for x in ("a", "b", "c", "e")]), encoding="utf-8")
    (ddir / "ambiguous.xes").write_text(
        _log([_event_single("a"), _event_single("b"),
              _event_ambiguous(["c", "d"]), _event_single("e")]), encoding="utf-8")
    # null-model baseline log: the ambiguity already resolved (here to d, the
    # wrong pick, so baseline accuracy is measurably below 1)
    (ddir / "baseline.xes").write_text(
        _log([_event_single(x) for x in ("a", "b", "d", "e")]), encoding="utf-8")
    return tmp_path


def _sheet_pairs(path):
    wb = openpyxl.load_workbook(path, read_only=True)
    return wb.sheetnames


def test_three_workbooks_per_dataset(data_dir, tmp_path):
    out = tmp_path / "results"
    run_all.build(DATASET, data_dir, out, repeats=1, limit=None)

    produced = sorted(p.name for p in out.glob("*.xlsx"))
    # the fixed-bound workbook carries the maxK in its name (default maxK=3)
    assert produced == ["SH_adaptive.xlsx", "SH_baseline.xlsx", "SH_maxk_3.xlsx"]
    assert list(out.glob("*.csv")) == []  # CSVs are gone


def test_sheets_and_summary(data_dir, tmp_path):
    out = tmp_path / "results"
    run_all.build(DATASET, data_dir, out, repeats=1, limit=None)

    for algo, fname in (("maxk", "SH_maxk_3.xlsx"), ("adaptive", "SH_adaptive.xlsx"),
                        ("baseline", "SH_baseline.xlsx")):
        path = out / fname
        names = _sheet_pairs(path)
        assert names[0] == "Summary"  # Summary first
        for s in ("AllEvents", "Results", "Fitness", "Runtime_GoldStandard"):
            assert s in names, f"{algo}: missing sheet {s}"

        wb = openpyxl.load_workbook(path, read_only=True)
        kv = {row[0]: row[1] for row in wb["Summary"].iter_rows(values_only=True)
              if row and row[0] is not None}
        acc = kv.get("All-events accuracy")
        assert isinstance(acc, (int, float)) and 0.0 <= acc <= 1.0


def test_maxk_only_run_names_and_records_the_bound(data_dir, tmp_path):
    """--algos maxk with --maxk N produces only the fixed-bound workbook, named
    with the maxk_N suffix, and records the bound in its Summary sheet."""
    out = tmp_path / "results"
    run_all.build(DATASET, data_dir, out, repeats=1, limit=None, maxk=5, algos=("maxk",))

    produced = sorted(p.name for p in out.glob("*.xlsx"))
    assert produced == ["SH_maxk_5.xlsx"]

    wb = openpyxl.load_workbook(out / "SH_maxk_5.xlsx", read_only=True)
    kv = {r[0]: r[1] for r in wb["Summary"].iter_rows(values_only=True)
          if r and r[0] is not None}
    assert kv["maxK"] == 5


def test_baseline_accuracy_below_one(data_dir, tmp_path):
    """The baseline log resolved the ambiguity to the wrong label, so its
    all-events accuracy must be < 1 while the adaptive one disambiguates to 1."""
    out = tmp_path / "results"
    run_all.build(DATASET, data_dir, out, repeats=1, limit=None)

    def acc(algo):
        wb = openpyxl.load_workbook(out / f"SH_{algo}.xlsx", read_only=True)
        kv = {r[0]: r[1] for r in wb["Summary"].iter_rows(values_only=True)
              if r and r[0] is not None}
        return kv["All-events accuracy"]

    assert acc("baseline") < 1.0
    assert acc("adaptive") == 1.0
